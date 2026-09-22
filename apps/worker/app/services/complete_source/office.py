"""Bounded Office source readers. No macros, external links, or conversion service."""

from __future__ import annotations

import io
import posixpath
import warnings
import zipfile
from datetime import date, datetime, time
from pathlib import PurePosixPath

import lxml.etree as etree

from .export import CompleteSourceError, ExportBuilder, MAX_UNITS, json_text

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "v": "urn:schemas-microsoft-com:vml",
}
_PICTURE_URI = "http://schemas.openxmlformats.org/drawingml/2006/picture"
# Plain text-box geometry only. Flowchart/callout/arrow presets carry meaning of
# their own, so shapes with any other geometry still fail closed.
_TEXT_BOX_PRESETS = {"rect", "textBox"}


def _text_box_content(node: etree._Element) -> etree._Element:
    """Return the `w:txbxContent` of a plain text box, or fail closed.

    Accepts a DrawingML `wps:wsp` whose geometry is a rectangle/text box, or a
    VML `v:rect`/`v:shape` text box. Geometry-only shapes, connectors, groups and
    any other preset are not silently flattened to their text.
    """
    local = etree.QName(node).localname
    if etree.QName(node).namespace == NS["wps"] and local == "wsp":
        presets = node.xpath("./wps:spPr/a:prstGeom/@prst", namespaces=NS)
        custom = node.xpath("./wps:spPr/a:custGeom", namespaces=NS)
        contents = node.xpath("./wps:txbx/w:txbxContent", namespaces=NS)
    elif etree.QName(node).namespace == NS["v"] and local in {"rect", "shape"}:
        presets = [] if local == "rect" else [node.get("type", "")]
        custom = [] if local == "rect" or node.get("type", "").endswith("_t202") else [node]
        contents = node.xpath("./v:textbox/w:txbxContent", namespaces=NS)
    else:
        raise CompleteSourceError("complete-source-unsupported-embedded-content")
    if custom or len(contents) != 1 or any(
        preset not in _TEXT_BOX_PRESETS and not preset.endswith("_t202") for preset in presets
    ):
        raise CompleteSourceError("complete-source-unsupported-embedded-content")
    return contents[0]


def safe_office_zip(data: bytes) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        entries = archive.infolist()
        if len(entries) > 2048 or sum(info.file_size for info in entries) > 64 * 1024 * 1024:
            raise CompleteSourceError("complete-source-resource-limit")
        names: set[str] = set()
        for info in entries:
            path = PurePosixPath(info.filename)
            if info.filename in names or "\\" in info.filename or path.is_absolute() or ".." in path.parts or info.flag_bits & 1:
                raise CompleteSourceError("complete-source-archive-invalid")
            names.add(info.filename)
            if info.file_size > 16 * 1024 * 1024 or (info.file_size >= 1024 * 1024 and info.file_size / max(1, info.compress_size) > 100):
                raise CompleteSourceError("complete-source-resource-limit")
            if "vbaproject" in info.filename.lower() or "externallinks/" in info.filename.lower():
                raise CompleteSourceError("complete-source-unsupported-embedded-content")
        return archive
    except (zipfile.BadZipFile, OSError) as exc:
        raise CompleteSourceError("complete-source-archive-invalid") from exc


def xml(data: bytes):
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise CompleteSourceError("complete-source-archive-invalid")
    return etree.fromstring(data, etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))


def export_xlsx(builder: ExportBuilder, data: bytes) -> list[str]:
    import openpyxl

    with safe_office_zip(data) as archive:
        if "xl/workbook.xml" not in archive.namelist():
            raise CompleteSourceError("complete-source-signature-invalid")
        # Validate every XML member before openpyxl can inflate/interpret it.
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                xml(archive.read(name))
        if any(name.startswith("xl/charts/") for name in archive.namelist()):
            raise CompleteSourceError("complete-source-unsupported-embedded-content")
    try:
        with warnings.catch_warnings():
            # openpyxl warns when it drops unsupported drawings/extensions. That
            # cannot become a successful complete export.
            warnings.simplefilter("error", UserWarning)
            workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
            cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
        units = list(workbook.sheetnames)
        if not 1 <= len(units) <= MAX_UNITS:
            raise CompleteSourceError("complete-source-resource-limit")
        total_cells = 0
        for sheet in workbook.worksheets:
            if sheet.max_row * sheet.max_column > 100_000:
                raise CompleteSourceError("complete-source-resource-limit")
            cells = []
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value
                    if isinstance(value, (datetime, date, time)):
                        value = value.isoformat()
                    if not isinstance(value, (str, int, float, bool)):
                        raise CompleteSourceError("complete-source-unsupported-cell")
                    record = {"coordinate": cell.coordinate, "type": cell.data_type, "value": value, "number_format": cell.number_format}
                    if cell.comment is not None:
                        record["comment"] = cell.comment.text
                    if cell.hyperlink is not None:
                        record["hyperlink"] = cell.hyperlink.target or cell.hyperlink.location
                    if cell.data_type == "f":
                        cached_value = cached[sheet.title][cell.coordinate].value
                        if isinstance(cached_value, (datetime, date, time)):
                            cached_value = cached_value.isoformat()
                        record["cached_value"] = cached_value
                    cells.append(record)
            total_cells += len(cells)
            builder.text(sheet.title, json_text({
                "format": "workbook-cells@v1", "sheet": sheet.title,
                "state": sheet.sheet_state, "rows": sheet.max_row, "columns": sheet.max_column,
                "merged_ranges": [str(item) for item in sheet.merged_cells.ranges], "cells": cells,
            }))
            for image in getattr(sheet, "_images", []):
                anchor = image.anchor
                marker = getattr(anchor, "_from", None)
                builder.text(sheet.title, json_text({"image_anchor": {"row": getattr(marker, "row", 0) + 1, "column": getattr(marker, "col", 0) + 1}}))
                builder.image(sheet.title, image._data())
        workbook.close()
        cached.close()
        if total_cells == 0 and not builder.resources:
            raise CompleteSourceError("complete-source-empty")
        return units
    except CompleteSourceError:
        raise
    except Exception as exc:
        raise CompleteSourceError("complete-source-workbook-invalid") from exc


def export_docx(builder: ExportBuilder, data: bytes) -> None:
    try:
        with safe_office_zip(data) as archive:
            names = archive.namelist()
            if "word/document.xml" not in names:
                raise CompleteSourceError("complete-source-signature-invalid")
            if any(name.startswith(("word/charts/", "word/embeddings/")) for name in names):
                raise CompleteSourceError("complete-source-unsupported-embedded-content")
            for name in names:
                if name.endswith((".xml", ".rels")):
                    xml(archive.read(name))
            # Body order is preserved; named ancillary parts are not silently dropped.
            content_names = ["word/document.xml"] + sorted(name for name in names if name.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes")) and name.endswith(".xml"))
            has_body_text = False
            for name in content_names:
                root = xml(archive.read(name))
                rel_name = posixpath.join(posixpath.dirname(name), "_rels", posixpath.basename(name) + ".rels")
                relationships = {}
                if rel_name in names:
                    for rel in xml(archive.read(rel_name)):
                        relationships[rel.get("Id")] = rel
                if name != "word/document.xml":
                    builder.text("1", json_text({"document_part": name.removeprefix("word/")}))
                buffer: list[str] = []

                def flush() -> None:
                    builder.text("1", "".join(buffer))
                    buffer.clear()

                def walk(node) -> None:
                    nonlocal has_body_text
                    local = etree.QName(node).localname
                    namespace = etree.QName(node).namespace
                    if namespace == NS["w"] and local == "sym":
                        # w:char is a font-specific glyph code, not trustworthy
                        # Unicode. Without a font renderer it cannot be dropped
                        # or guessed (e.g. a checked Wingdings box).
                        raise CompleteSourceError("complete-source-unsupported-symbol")
                    if namespace == NS["m"] and local == "oMath":
                        has_body_text = has_body_text or bool("".join(node.itertext()).strip())
                        buffer.append(json_text({"format": "word-math-xml", "expression": etree.tostring(node, encoding="unicode")}))
                        return
                    if local in {"altChunk", "OLEObject", "chart"}:
                        raise CompleteSourceError("complete-source-unsupported-embedded-content")
                    if local == "graphicData" and node.get("uri") == NS["wps"]:
                        # A plain text box (e.g. the page-number footer Word generates)
                        # is ordinary paragraph content; export it in place and marked.
                        children = list(node)
                        if len(children) != 1 or etree.QName(children[0]).localname != "wsp":
                            raise CompleteSourceError("complete-source-unsupported-embedded-content")
                        buffer.append("\n[TEXTBOX]\n")
                        for child in _text_box_content(children[0]):
                            walk(child)
                        buffer.append("\n")
                        return
                    if local == "graphicData" and node.get("uri") != _PICTURE_URI:
                        raise CompleteSourceError("complete-source-unsupported-embedded-content")
                    if local == "pict" and not node.xpath(".//*[local-name()='imagedata']"):
                        # Legacy VML: exactly one text box shape (plus its shapetype
                        # definition), nothing else drawn or embedded.
                        children = [child for child in node if etree.QName(child).localname != "shapetype"]
                        if len(children) != 1:
                            raise CompleteSourceError("complete-source-unsupported-embedded-content")
                        buffer.append("\n[TEXTBOX]\n")
                        for child in _text_box_content(children[0]):
                            walk(child)
                        buffer.append("\n")
                        return
                    if local == "AlternateContent":
                        # Choice/Fallback represent the same authored object, not two
                        # independent images. Unsupported choice still fails closed.
                        selected = next(iter(node), None)
                        if selected is not None:
                            walk(selected)
                        return
                    if local in {"blip", "imagedata"}:
                        rel_id = node.get(f"{{{NS['r']}}}embed") or node.get(f"{{{NS['r']}}}id")
                        rel = relationships.get(rel_id)
                        if rel is None or rel.get("TargetMode") == "External":
                            raise CompleteSourceError("complete-source-external-resource-required")
                        target = posixpath.normpath(posixpath.join(posixpath.dirname(name), rel.get("Target", "")))
                        if target not in names or not target.startswith("word/media/"):
                            raise CompleteSourceError("complete-source-partial")
                        flush()
                        builder.image("1", archive.read(target))
                        return
                    if namespace == NS["w"] and local == "t":
                        has_body_text = has_body_text or bool((node.text or "").strip())
                        buffer.append(node.text or "")
                    elif local == "tab":
                        buffer.append("\t")
                    elif local in {"br", "cr"}:
                        buffer.append("\n")
                    elif local in {"gridSpan", "vMerge"}:
                        buffer.append(json_text({local: node.get(f"{{{NS['w']}}}val", "continue")}))
                    elif local == "tc":
                        buffer.append("\n[CELL]\n")
                    elif local == "tr":
                        buffer.append("\n[ROW]\n")
                    for child in node:
                        walk(child)
                    if local in {"p", "tc", "tr", "tbl"}:
                        buffer.append("\n")

                walk(root)
                flush()
            if not has_body_text and not builder.resources:
                raise CompleteSourceError("complete-source-empty")
    except CompleteSourceError:
        raise
    except Exception as exc:
        raise CompleteSourceError("complete-source-docx-invalid") from exc
