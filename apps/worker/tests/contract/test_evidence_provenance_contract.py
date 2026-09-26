"""Database-free parser -> chunk JSON -> hydration provenance contract."""
from __future__ import annotations

import json
import os
from pathlib import Path

os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://test:test@localhost/test")
os.environ.setdefault("TMP_PATH", "/tmp/knowhere-test")
os.environ.setdefault("S3_BUCKET_NAME", "test-uploads")
os.environ.setdefault("S3_ACCESS_KEY_ID", "test")
os.environ.setdefault("S3_SECRET_ACCESS_KEY", "test")
os.environ.setdefault("S3_TEMP_PATH", "/tmp")

import pytest
from docx import Document
from app.services.document_parser.formats.docx.image_summary_scheduler import (
    DocxImageOccurrence, DocxImageSummaryScheduler,
)
from app.services.document_parser.formats.docx.parser import convert_doc2dics, _row_values_to_parsed_row
from app.services.document_parser.formats.docx.table_html import table2html
from app.services.document_parser.support.parser_rows import ParsedRow, PARSER_ROW_COLUMNS
from shared.services.ai.summary.model import AssetSummary
from shared.services.chunks.dataframe_chunk_converter import dataframe_to_chunks
from shared.services.chunks.evidence_provenance import (
    decode_provenance, encode_provenance, join_text, marked, text_metadata,
)
from shared.services.retrieval.hydration import evidence_compose as renderer


def test_scheduler_chunking_json_roundtrip_does_not_guess_equal_prose(tmp_path):
    description = "same authored and generated 文😀"
    holder = [marked(description, "source"), marked("x" * 40000, "source"), "placeholder"]
    row = ParsedRow(content="placeholder", path="images/a.png", type="image", know_id="img", addtime="now").to_list()
    scheduler = DocxImageSummaryScheduler(should_summarize=True)
    scheduler._results_by_hash["hash"] = AssetSummary(summary=description, title="")
    scheduler.register_occurrence(DocxImageOccurrence("hash", "image-1", row, holder, 2))
    frame = convert_doc2dics(
        {"heading": "Heading", "content": holder}, [row], str(tmp_path),
        {"doc_type": "document", "summary_txt": False, "stopwords": []},
        relative_root="doc",
    )
    chunks = json.loads(json.dumps(dataframe_to_chunks(frame)))
    parts = []
    for chunk in chunks:
        # This is the existing publication JSON metadata shape, not a DB-write proof.
        restored = {"chunk_id": chunk["chunk_id"], "document_id": "doc", "chunk_type": "text",
                    "content": chunk["content"], "chunk_metadata": chunk["metadata"]}
        restored["composed"] = renderer.compose_evidence_parts(restored, {})
        parts.extend(renderer.collect_evidence([restored]))
    assert any(p["text"] == description and p["provenance"] == "source" for p in parts)
    assert any(p["text"] == description and p["provenance"] == "generated-image-description" for p in parts)
    assert all(p["document_id"] == "doc" and p["chunk_id"] for p in parts)
    assert all(p["provenance"] != "unknown" for p in parts)
    assert len(frame) >= 3  # image plus divided body, not an unsplit toy sample
    assert _row_values_to_parsed_row(row).extra_metadata == row[PARSER_ROW_COLUMNS.index("extra_metadata")]
    assert _row_values_to_parsed_row(row[:11]).extra_metadata == {}


def test_text_adapter_provenance_survives_real_markdown_parser(tmp_path, monkeypatch):
    from app.services.document_parser.formats.markdown import parser

    monkeypatch.setattr(parser, "tokenize2stw_remove", lambda *args: [])
    for trusted in (True, False):
        frame = parser.parse_md(
            str(tmp_path / str(trusted)), source_type="md",
            lines_with_heading=["# Heading", "authored sentence", "second sentence"],
            base_llm_paras={"summary_txt": False, "summary_table": False, "summary_image": False, "stopwords": []},
            relative_root="document", plain_text_source=trusted,
        )
        chunks = dataframe_to_chunks(frame)
        assert chunks
        text = chunks[0]["content"]
        decoded = decode_provenance(text, chunks[0]["metadata"]["evidence_provenance"])
        assert ("authored sentence", "source" if trusted else "unknown") in decoded.segments


def test_table_html_binding_and_inline_hydration(monkeypatch):
    document = Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "identical description"
    html = table2html(table, {(0, 0): marked("identical description", "generated-image-description")})
    target = {"chunk_id": "table", "chunk_type": "table", "chunk_metadata": {
        "table_evidence_provenance": encode_provenance(html)}}
    body = join_text([marked("原文", "source"), marked("[tables/t.html]", "system")])
    row = {"chunk_type": "text", "content": str(body), "chunk_metadata": {
        **text_metadata(body), "connect_to": [{"target": "table", "relation": "embeds", "ref": "[tables/t.html]"}]}}
    monkeypatch.setattr(renderer, "load_table_html", lambda _: str(html))
    parts = renderer.compose_evidence_parts(row, {"table": target})
    matches = [p["provenance"] for p in parts if p.get("text") == "identical description"]
    assert matches == ["source", "generated-image-description"]
    assert "<table" in renderer.flatten_parts(parts)
    target["chunk_metadata"]["table_evidence_provenance"]["sha256"] = "stale"
    parts = renderer.compose_evidence_parts(row, {"table": target})
    assert any(p["text"] == html and p["provenance"] == "unknown" for p in parts)


@pytest.mark.parametrize("embedded", [False, True])
@pytest.mark.parametrize("trusted", [True, False])
def test_markdown_assets_parser_to_hydration(tmp_path, monkeypatch, trusted, embedded):
    from PIL import Image
    from app.services.document_parser.formats.markdown import parser, deferred_summary, image_asset

    (tmp_path / "images").mkdir()
    Image.effect_noise((220, 220), 100).convert("RGB").save(tmp_path / "images/input.png")
    monkeypatch.setattr(image_asset, "discard_undersized_image_file", lambda *a, **kw: False)
    monkeypatch.setattr(parser, "tokenize2stw_remove", lambda *a: [])
    monkeypatch.setattr(deferred_summary, "summarize", lambda **kw: AssetSummary(title="renamed", summary="同名😀描述"))
    frame = parser.parse_md(
        str(tmp_path), source_type="md", plain_text_source=trusted,
        lines_with_heading=(["# Heading", "同名😀描述",
                             '<table><tr><td>文😀</td><td><img src="images/input.png"></td></tr></table>']
                            if embedded else ["# Heading", "同名😀描述", "![caption](images/input.png)",
                            "| 名称 | 数值 |", "| --- | --- |", "| 文😀 | 42 |"]),
        base_llm_paras={"summary_txt": False, "summary_table": True, "summary_image": True, "stopwords": []},
        relative_root="document",
    )
    chunks = json.loads(json.dumps(dataframe_to_chunks(frame)))
    rows = {c["chunk_id"]: {"chunk_id": c["chunk_id"], "document_id": "doc",
            "chunk_type": c["type"], "content": c["content"], "chunk_metadata": c["metadata"]} for c in chunks}
    monkeypatch.setattr(renderer, "load_table_html", lambda row: (tmp_path / row["chunk_metadata"]["file_path"]).read_text())
    monkeypatch.setattr(renderer, "_try_read_image", lambda row: {"type": "image", "provenance": "source"})
    parts = []
    for row in rows.values():
        row["composed"] = renderer.compose_evidence_parts(row, rows)
        parts.extend(renderer.collect_evidence([row]))
    expected = "source" if trusted else "unknown"
    assert any(p.get("text") == "同名😀描述" and p["provenance"] == expected for p in parts)
    assert any(p.get("text") == "同名😀描述" and p["provenance"] == "generated-image-description" for p in parts)
    assert any("文😀" in p.get("text", "") and p["provenance"] == expected for p in parts)
    assert any("<table" in p.get("text", "") and p["provenance"] == "system" for p in parts)
    if trusted:
        assert all(p["provenance"] != "unknown" for p in parts)
    assert all(p["document_id"] == "doc" for p in parts)


def test_table_serializer_preserves_exact_unicode_and_unterminated_entities():
    from app.services.document_parser.tables.html_provenance import annotate_table_html

    html = '<table><td>文😀 &amp; &amp x &#65; &#65 z</td></table>'
    annotated = annotate_table_html(html, text_kind="source")
    assert str(annotated) == html
    assert decode_provenance(html, encode_provenance(annotated)).segments == annotated.segments


def test_excel_real_workbook_asset_binding(tmp_path, monkeypatch):
    from openpyxl import Workbook
    from app.services.document_parser.formats.excel import table_parser
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["名称", "数值"])
    sheet.append(["原文😀 & <tag>", 42])
    source = tmp_path / "input.xlsx"
    workbook.save(source)
    monkeypatch.setattr(table_parser, "tokenize2stw_remove", lambda *a: [])
    frame = table_parser.parse_xlsx(str(source), "input.xlsx", str(tmp_path), "",
                                   base_llm_paras={"summary_table": False, "stopwords": []})
    chunks = json.loads(json.dumps(dataframe_to_chunks(frame)))
    assert chunks
    for chunk in chunks:
        row = {"chunk_type": "table", "content": chunk["content"], "chunk_metadata": chunk["metadata"]}
        monkeypatch.setattr(renderer, "load_table_html", lambda row: (tmp_path / row["chunk_metadata"]["file_path"]).read_text())
        parts = renderer.compose_evidence_parts(row, {})
        assert any("原文😀" in p["text"] and p["provenance"] == "source" for p in parts)
        assert any("<table" in p["text"] and p["provenance"] == "system" for p in parts)
        assert all(p["provenance"] != "unknown" for p in parts)
        row["chunk_metadata"]["table_evidence_provenance"]["sha256"] = "stale"
        assert all(p["provenance"] == "unknown" for p in renderer.compose_evidence_parts(row, {}))


@pytest.mark.parametrize("corruption", ["missing", "version", "hash", "gap", "overlap", "kind", "bool", "tail"])
def test_invalid_binding_is_wholly_unknown(corruption):
    text = join_text([marked("文😀", "source"), marked("generated", "generated-image-description")])
    binding = encode_provenance(text)
    if corruption == "missing": binding = None
    elif corruption == "version": binding["version"] = "future"
    elif corruption == "hash": binding["sha256"] = "wrong"
    elif corruption == "gap": binding["spans"][1]["start"] += 1
    elif corruption == "overlap": binding["spans"][1]["start"] -= 1
    elif corruption == "kind": binding["spans"][0]["provenance"] = "original"
    elif corruption == "bool": binding["spans"][0]["start"] = False
    elif corruption == "tail": binding["spans"].pop()
    assert decode_provenance(str(text), binding).segments == ((str(text), "unknown"),)


@pytest.mark.parametrize("branch,expected", [("figure", "generated-image-description"), ("text", "transcription"), ("disabled", "system"), ("empty", "system")])
def test_real_standalone_image_parser_branches(tmp_path, monkeypatch, branch, expected):
    from PIL import Image
    from app.services.document_parser.formats.image import parser
    from app.services.document_parser.assets import image_size_filter
    image = tmp_path / "input.png"
    Image.new("RGB", (100, 100)).save(image)
    monkeypatch.setattr(image_size_filter, "discard_undersized_image_file", lambda *a, **kw: False)
    monkeypatch.setattr(parser, "_get_vision_client", lambda: object())
    monkeypatch.setattr(parser, "ask_image", lambda *a, **kw: {"answer": branch})
    monkeypatch.setattr(parser, "transcribe", lambda **kw: "visual transcription")
    monkeypatch.setattr(parser, "summarize", lambda **kw: AssetSummary(summary="" if branch == "empty" else "generated description", title=""))
    frame = parser.parse_image(str(image), filename="input.png", output_dir=str(tmp_path / "out"),
                               base_llm_paras={"frag_desc": "", "summary_image": branch != "disabled"}, auto_rename=False)
    chunk = json.loads(json.dumps(dataframe_to_chunks(frame)))[0]
    monkeypatch.setattr(renderer, "_try_read_image", lambda _: None)
    parts = renderer.compose_evidence_parts({"chunk_type": "image", "content": chunk["content"], "chunk_metadata": chunk["metadata"]}, {})
    assert parts[-1]["provenance"] == expected
    assert parts[0]["provenance"] == "system"
    assert (tmp_path / "out/images/input.png").exists()
