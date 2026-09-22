from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest
from PIL import Image

from app.services.complete_source.export import CompleteSourceError, build_complete_source, digest
from shared.models.schemas.complete_source import COMPLETE_SOURCE_EXTENSIONS, CompleteSourceRequest


def png() -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (16, 12), "red").save(output, format="PNG")
    return output.getvalue()


def export(tmp_path: Path, data: bytes, name: str, kind: str, encoding: str | None = None):
    source = tmp_path / name
    source.write_bytes(data)
    request = CompleteSourceRequest.model_validate({
        "schema_version": "knowhere-complete-source-request@v1", "kind": kind,
        **({"encoding": encoding} if encoding else {}),
    })
    target = tmp_path / "result.zip"
    result, checksum, size = build_complete_source(
        job_id="job_complete_test", source_path=source, request=request, output_path=target,
    )
    raw = target.read_bytes()
    assert digest(raw) == checksum
    assert len(raw) == size
    assert result.input_sha256 == digest(data)
    assert result.input_byte_length == len(data)
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        assert set(archive.namelist()) == {"complete-source.json", *(resource.path for resource in result.resources)}
        assert json.loads(archive.read("complete-source.json")) == result.model_dump()
        for resource in result.resources:
            content = archive.read(resource.path)
            assert len(content) == resource.byte_length
            assert digest(content) == resource.sha256
        for part in result.parts:
            if part.type == "text":
                assert digest(part.text.encode("utf-8")) == part.sha256
    assert result.coverage.vision_required == bool(result.resources)
    return result, raw


@pytest.mark.parametrize("extension", sorted(COMPLETE_SOURCE_EXTENSIONS["code"]))
def test_code_keeps_exact_whitespace(tmp_path: Path, extension: str) -> None:
    text = "def example():\r\n\treturn 'a  b'\r\n\r\n# 尾部唯一标记\n"
    result, _ = export(tmp_path, text.encode(), "source" + extension, "code")
    assert result.parts[0].text == text
    assert result.coverage.text_extraction == "decoded"


def test_bom_and_explicit_encoding(tmp_path: Path) -> None:
    text = "空白  preserved\n  indent\n"
    result, _ = export(tmp_path, text.encode("utf-16"), "source.txt", "text")
    assert result.parts[0].text == text


def test_explicit_big5_is_admitted(tmp_path: Path) -> None:
    text = "繁體中文  保留空白"
    result, _ = export(tmp_path, text.encode("big5"), "source.txt", "text", "big5")
    assert result.parts[0].text == text
    result, _ = export(tmp_path, text.encode("gb18030"), "source.txt", "text", "gb18030")
    assert result.parts[0].text == text


def test_markdown_keeps_code_and_blank_lines(tmp_path: Path) -> None:
    text = "# Title\n\n```py\n    a = 1  + 2\n```\n\nTAIL\n"
    result, _ = export(tmp_path, text.encode(), "source.md", "document")
    assert result.parts[0].text == text


@pytest.mark.parametrize("name,text", [("table.csv", '"head","tail"\r\n"quoted\nnewline","end, cell"'), ("table.csv", "h1;h2\n1;last"), ("table.tsv", "h1\th2\n1\tlast")])
def test_delimited_cells_and_tail(tmp_path: Path, name: str, text: str) -> None:
    result, _ = export(tmp_path, text.encode(), name, "table")
    rows = json.loads(result.parts[0].text)["rows"]
    assert len(rows) == 2
    assert rows[-1][-1]["value"] in {"last", "end, cell"}
    if "quoted" in text:
        assert rows[-1][0]["value"] == "quoted\nnewline"


def test_pdf_preserves_toc_tail_and_scan_page(tmp_path: Path) -> None:
    import pymupdf

    document = pymupdf.open()
    document.new_page().insert_text((30, 30), "TOC_ONLY_DETAIL")
    scan = document.new_page()
    scan.insert_image(pymupdf.Rect(20, 20, 100, 100), stream=png())
    document.new_page().insert_text((30, 30), "TAIL_PAGE_UNIQUE_DETAIL")
    data = document.tobytes()
    document.close()
    result, _ = export(tmp_path, data, "source.pdf", "document")
    assert result.coverage.covered_units == ["1", "2", "3"]
    assert result.coverage.total_units == 3
    assert result.coverage.text_extraction == "mixed"
    assert [part.unit for part in result.parts if part.type == "image"] == ["1", "2", "3"]
    assert "TOC_ONLY_DETAIL" in result.parts[0].text
    assert any(part.type == "text" and "TAIL_PAGE_UNIQUE_DETAIL" in part.text for part in result.parts)


def test_short_scanned_pdf_is_visual_not_fake_ocr(tmp_path: Path) -> None:
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_image(pymupdf.Rect(20, 20, 100, 100), stream=png())
    result, _ = export(tmp_path, document.tobytes(), "source.pdf", "document")
    document.close()
    assert result.coverage.text_extraction == "unavailable"
    assert len(result.parts) == 1 and result.parts[0].type == "image"


@pytest.mark.parametrize("format,name", [("PNG", "source.png"), ("JPEG", "source.jpg"), ("WEBP", "source.webp")])
def test_images_are_real_pixels_without_summary(tmp_path: Path, format: str, name: str) -> None:
    output = io.BytesIO()
    Image.new("RGB", (30, 20), "blue").save(output, format=format)
    result, _ = export(tmp_path, output.getvalue(), name, "figure")
    assert len(result.parts) == 1 and result.parts[0].type == "image"
    assert result.resources[0].width == 30
    assert result.resources[0].height == 20
    assert result.resources[0].sha256 == digest(output.getvalue())


def test_xlsx_keeps_types_formula_hidden_tail_sheet_and_images(tmp_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as WorkbookImage

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"], sheet["B1"], sheet["A2"], sheet["B2"] = "heading", 42, "tail  cell", "=B1*2"
    sheet.merge_cells("A4:B4")
    sheet.add_image(WorkbookImage(io.BytesIO(png())), "C5")
    last = workbook.create_sheet("Hidden tail")
    last.sheet_state = "hidden"
    last["A1"] = "FINAL_SHEET_DETAIL"
    output = io.BytesIO()
    workbook.save(output)
    result, _ = export(tmp_path, output.getvalue(), "source.xlsx", "table")
    assert result.coverage.covered_units == ["Sheet", "Hidden tail"]
    payload = json.loads(result.parts[0].text)
    cells = {cell["coordinate"]: cell for cell in payload["cells"]}
    assert cells["B1"]["value"] == 42
    assert cells["B2"]["type"] == "f" and cells["B2"]["value"] == "=B1*2"
    assert payload["merged_ranges"] == ["A4:B4"]
    assert any(part.type == "text" and "FINAL_SHEET_DETAIL" in part.text for part in result.parts)
    assert len(result.resources) == 1


def test_docx_keeps_order_table_and_pixels(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("BEFORE  IMAGE")
    document.add_picture(io.BytesIO(png()))
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "FIRST_CELL"
    table.cell(1, 1).text = "LAST_CELL"
    document.add_paragraph("TAIL_DOCX")
    output = io.BytesIO()
    document.save(output)
    result, _ = export(tmp_path, output.getvalue(), "source.docx", "document")
    assert "BEFORE  IMAGE" in result.parts[0].text
    assert result.parts[1].type == "image"
    assert "FIRST_CELL" in result.parts[2].text and "LAST_CELL" in result.parts[2].text
    assert "TAIL_DOCX" in result.parts[2].text


def test_html_visible_pre_table_and_embedded_image(tmp_path: Path) -> None:
    import base64

    text = '<html><head><title>NOT_BODY</title></head><body><script>NOT_VISIBLE</script><pre>  a  b\n\tcode</pre><table><tr><td rowspan="2">cell</td></tr><tr><td>TAIL</td></tr></table><img src="data:image/png;base64,' + base64.b64encode(png()).decode() + '"></body></html>'
    result, _ = export(tmp_path, text.encode(), "source.html", "document")
    all_text = "".join(part.text for part in result.parts if part.type == "text")
    assert "NOT_BODY" not in all_text and "NOT_VISIBLE" not in all_text
    assert "  a  b\n\tcode" in all_text
    assert '"row_span":2' in all_text and "TAIL" in all_text
    assert len(result.resources) == 1


@pytest.mark.parametrize("name,text", [("source.bib", '@article{key, title={Full {nested} Title}, abstract={Last field}}'), ("source.ris", "TY  - JOUR\nTI  - Complete title\nAB  - TAIL_ABSTRACT\nER  - \n"), ("source.json", '[{"id":"a","title":"Complete title","abstract":"TAIL_ABSTRACT"}]'), ("source.txt", "Full authored abstract")])
def test_bibliography_keeps_semantic_fields(tmp_path: Path, name: str, text: str) -> None:
    result, _ = export(tmp_path, text.encode(), name, "literature")
    assert "title" in result.parts[0].text.lower() or "abstract" in result.parts[0].text.lower()


@pytest.mark.parametrize("name,kind,data,code", [
    ("x.png", "figure", b"filename only", "signature-invalid"),
    ("x.pdf", "document", b"not a pdf", "signature-invalid"),
    ("x.txt", "text", b"a\x00b", "text-invalid"),
    ("x.txt", "text", b"\xff\xff", "encoding-required"),
    ("x.txt", "text", b" \n\t", "empty"),
    ("x.csv", "table", b"one column\nonly", "table-invalid"),
    ("x.json", "literature", b'{"id":"only-id"}', "literature-invalid"),
    ("x.ris", "literature", b"TY  - JOUR\nTI  - title", "literature-invalid"),
    ("x.bib", "literature", b"@article{key, title={broken}", "literature-invalid"),
    ("x.html", "document", b'<img src="https://private.example/image.png">', "external-resource-required"),
    ("x.md", "document", b"![external](file.png)", "external-resource-required"),
    ("x.html", "document", b"<canvas>fake text</canvas>", "unsupported-embedded-content"),
])
def test_invalid_or_partial_input_never_publishes_bundle(tmp_path: Path, name: str, kind: str, data: bytes, code: str) -> None:
    with pytest.raises(CompleteSourceError, match="complete-source-" + code):
        export(tmp_path, data, name, kind)
    assert not (tmp_path / "result.zip").exists()


def test_archive_traversal_rejected_before_inflation(tmp_path: Path) -> None:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("../escaped.xml", "unsafe")
    with pytest.raises(CompleteSourceError, match="archive-invalid"):
        export(tmp_path, output.getvalue(), "source.docx", "document")


def test_deterministic_zip_metadata(tmp_path: Path) -> None:
    _, first = export(tmp_path, b"exact  text\n", "source.txt", "text")
    _, second = export(tmp_path, b"exact  text\n", "source.txt", "text")
    assert first == second


def test_bibtex_parses_parentheses_macro_and_rejects_later_id_only_entry(tmp_path: Path) -> None:
    text = '@string{venue="Complete Journal"}\n@article(key, title={Full {nested} title}, journal=venue)'
    result, _ = export(tmp_path, text.encode(), "source.bib", "literature")
    assert result.parts[0].text == text
    with pytest.raises(CompleteSourceError, match="literature-invalid"):
        export(tmp_path, (text + "\n@article{bare,year={2026}}").encode(), "source.bib", "literature")


def test_markdown_embedded_pixels_are_parts_not_base64_prose(tmp_path: Path) -> None:
    import base64

    encoded = base64.b64encode(png()).decode()
    text = "before\n![caption](data:image/png;base64," + encoded + ")\nTAIL"
    result, _ = export(tmp_path, text.encode(), "source.md", "document")
    assert [part.type for part in result.parts] == ["text", "text", "image", "text"]
    assert all(encoded not in part.text for part in result.parts if part.type == "text")


def test_docx_missing_image_fails_instead_of_returning_only_text(tmp_path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("BODY")
    document.add_picture(io.BytesIO(png()))
    output = io.BytesIO()
    document.save(output)
    broken = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as original, zipfile.ZipFile(broken, "w") as target:
        for name in original.namelist():
            if not name.startswith("word/media/"):
                target.writestr(name, original.read(name))
    with pytest.raises(CompleteSourceError, match="partial"):
        export(tmp_path, broken.getvalue(), "source.docx", "document")


def test_output_budget_blocks_before_zip_without_truncation(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    from app.services.complete_source import export as module

    monkeypatch.setattr(module, "MAX_EXPORT_BYTES", 10)
    with pytest.raises(CompleteSourceError, match="resource-limit"):
        export(tmp_path, b"a" * 11, "source.txt", "text")
    assert not (tmp_path / "result.zip").exists()


@pytest.mark.parametrize("extension", ["docx", "xlsx"])
def test_empty_office_metadata_is_not_source_content(tmp_path: Path, extension: str) -> None:
    from docx import Document
    from openpyxl import Workbook

    output = io.BytesIO()
    if extension == "docx":
        Document().save(output)
    else:
        Workbook().save(output)
    with pytest.raises(CompleteSourceError, match="complete-source-empty"):
        export(tmp_path, output.getvalue(), "source." + extension, "document" if extension == "docx" else "table")


def test_html_table_retains_caption_and_mathml_structure(tmp_path: Path) -> None:
    text = '<table><caption>CAPTION_SENTINEL</caption><tr><td>before<math xmlns="http://www.w3.org/1998/Math/MathML"><mfrac><mi>a</mi><mi>b</mi></mfrac></math>after</td></tr></table>'
    result, _ = export(tmp_path, text.encode(), "source.html", "table")
    parts = "".join(part.text for part in result.parts if part.type == "text")
    assert "CAPTION_SENTINEL" in parts
    assert '"format":"mathml"' in parts
    assert "<mfrac><mi>a</mi><mi>b</mi></mfrac>" in parts
    assert parts.index("before") < parts.index("<mfrac>") < parts.index("after")


@pytest.mark.parametrize("tag", ["svg", "canvas", "object"])
@pytest.mark.parametrize("location", ["cell", "caption", "table-child"])
def test_html_table_rejects_unsupported_descendants(tmp_path: Path, tag: str, location: str) -> None:
    embedded = f"<{tag}>visual fallback</{tag}>"
    content = {
        "cell": f"<tr><td>{embedded}</td></tr>",
        "caption": f"<caption>{embedded}</caption><tr><td>body</td></tr>",
        "table-child": f"{embedded}<tr><td>body</td></tr>",
    }[location]
    with pytest.raises(CompleteSourceError, match="unsupported-embedded-content"):
        export(tmp_path, ("<table>" + content + "</table>").encode(), "source.html", "table")
    assert not (tmp_path / "result.zip").exists()


def test_docx_font_encoded_symbol_cannot_be_silently_dropped(tmp_path: Path) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    paragraph = document.add_paragraph("Status: ")
    symbol = OxmlElement("w:sym")
    symbol.set(qn("w:font"), "Wingdings")
    symbol.set(qn("w:char"), "F0FE")
    paragraph.add_run()._r.append(symbol)
    output = io.BytesIO()
    document.save(output)
    with pytest.raises(CompleteSourceError, match="complete-source-unsupported-symbol"):
        export(tmp_path, output.getvalue(), "source.docx", "document")
    assert not (tmp_path / "result.zip").exists()


_WPS_TEXT_BOX = (
    '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
    ' xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
    ' xmlns:v="urn:schemas-microsoft-com:vml">'
    '<mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
    '<wps:wsp><wps:spPr><a:prstGeom prst="{preset}"/></wps:spPr>'
    '<wps:txbx><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></wps:txbx></wps:wsp>'
    '</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
    '<mc:Fallback><w:pict><v:shapetype id="_x0000_t202"/><v:shape type="#_x0000_t202"><v:textbox><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>'
    '</mc:AlternateContent>'
)


def _docx_with_footer_shape(markup: str) -> bytes:
    from docx import Document
    from lxml import etree

    document = Document()
    document.add_paragraph("BODY_BEFORE_FOOTER")
    footer_paragraph = document.sections[0].footer.paragraphs[0]
    footer_paragraph.add_run("Page ")
    footer_paragraph.runs[0]._r.append(etree.fromstring(markup))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_plain_text_box_in_footer_exports_its_text_once(tmp_path: Path) -> None:
    data = _docx_with_footer_shape(_WPS_TEXT_BOX.format(preset="rect", text="FOOTER_BOX_SENTINEL"))
    result, _ = export(tmp_path, data, "source.docx", "document")
    all_text = "".join(part.text for part in result.parts if part.type == "text")
    assert "BODY_BEFORE_FOOTER" in all_text
    assert all_text.count("FOOTER_BOX_SENTINEL") == 1
    assert "[TEXTBOX]" in all_text
    assert all_text.index("BODY_BEFORE_FOOTER") < all_text.index("FOOTER_BOX_SENTINEL")


def test_docx_legacy_vml_text_box_exports_its_text(tmp_path: Path) -> None:
    markup = (
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml">'
        '<v:shapetype id="_x0000_t202"/><v:rect><v:textbox><w:txbxContent><w:p><w:r><w:t>VML_BOX_SENTINEL</w:t></w:r></w:p></w:txbxContent></v:textbox></v:rect></w:pict>'
    )
    result, _ = export(tmp_path, _docx_with_footer_shape(markup), "source.docx", "document")
    all_text = "".join(part.text for part in result.parts if part.type == "text")
    assert all_text.count("VML_BOX_SENTINEL") == 1


@pytest.mark.parametrize("markup", [
    _WPS_TEXT_BOX.format(preset="flowChartDecision", text="DECISION"),
    _WPS_TEXT_BOX.replace("<wps:txbx>", "<wps:txbx><w:txbxContent/>").format(preset="rect", text="TWO_BOXES"),
    '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml">'
    '<v:shape path="m0,0l100,100e"><v:textbox><w:txbxContent><w:p><w:r><w:t>CUSTOM_PATH</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>',
    '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml">'
    '<v:line from="0,0" to="1,1"/><v:rect><v:textbox><w:txbxContent><w:p><w:r><w:t>BOX_PLUS_LINE</w:t></w:r></w:p></w:txbxContent></v:textbox></v:rect></w:pict>',
])
def test_docx_non_text_box_shapes_still_fail_closed(tmp_path: Path, markup: str) -> None:
    with pytest.raises(CompleteSourceError, match="complete-source-unsupported-embedded-content"):
        export(tmp_path, _docx_with_footer_shape(markup), "source.docx", "document")
    assert not (tmp_path / "result.zip").exists()
